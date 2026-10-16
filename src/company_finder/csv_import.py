"""CSV / Excel şirket listesi içe aktarma ve elle şirket ekleme.

`import_companies()` bir `.csv` veya `.xlsx` dosyasını okur, her satırı doğrular ve geçerli
olanları veritabanı oturumuna ekler (COMMIT ETMEZ — `src/core/db.py`'deki "commit çağırana
aittir" kuralı burada da geçerlidir). `add_company()` aynı doğrulama çekirdeğini tek bir
şirket için kullanır (arayüzdeki "elle ekle" formu bunu çağıracak).

Sütun adları esnektir (büyük/küçük harf ve Türkçe karakter duyarsız, bkz. `COLUMN_ALIASES`).
Zorunlu alan: ad + (site veya e-posta). Alan adı (`domain`) önce siteden, yoksa e-postadan
`normalize_domain()` ile çıkarılır — bu fonksiyon artık sorgu/çapa/port/sondaki nokta gibi
farkları da tekilleştiriyor (bkz. `src/core/models.py`), ama bir DOĞRULAYICI değildir: "yok"
gibi anlamsız bir siteyi de aynen geçirir. Biçim doğrulaması burada yapılır.

Hiçbir satır tüm içe aktarmayı durdurmaz: bir satırdaki hata (bilinen ya da beklenmeyen)
sadece o satırı atlar, `ImportReport`'a Türkçe bir sebeple kaydedilir.
"""
import csv
import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterator, Optional

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlmodel import Session, select

from src.core.blocklist import blocked_domains
from src.core.models import Company, normalize_domain
from src.core.text import fold_turkish
from src.company_finder.email_extractor import is_role_address, normalize_email
from src.company_finder.free_mail import is_free_mail_domain
from src.llm.base import filter_tags

logger = logging.getLogger(__name__)

CSV_ENCODINGS: tuple[str, ...] = ("utf-8", "utf-8-sig", "cp1254")

# Excel'in "Unicode Text" (UTF-16) çıktısını erkenden yakalamak için: cp1254 tek baytlık
# olduğundan neredeyse her byte dizisini "başarıyla" (ama yanlış) çözer, bu yüzden UTF-16
# dosyaları CSV_ENCODINGS listesine hiç ulaşmadan burada elenir.
_UTF16_BOMS: tuple[bytes, ...] = (b"\xff\xfe", b"\xfe\xff")

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "name": ("name", "ad", "sirket", "firma", "sirket adi", "company"),
    "website": ("website", "web", "site", "url", "web sitesi"),
    "email": ("email", "e-posta", "eposta", "mail", "e posta"),
    "city": ("city", "sehir", "il"),
    "about_text": ("about", "hakkinda", "aciklama"),
    "tags": ("tags", "etiketler", "etiket"),
}

REASON_NO_NAME = "ad boş"
REASON_NO_CONTACT = "site veya e-posta gerekli"
REASON_FREE_MAIL_NEEDS_WEBSITE = "web sitesi gerekli"
REASON_BAD_DOMAIN = "alan adı çözülemedi"
REASON_BLOCKED = "alan adı engelli listesinde (bu şirkete bir daha yazılmaz)"
REASON_DUPLICATE_DB = "aynı alan adı zaten kayıtlı"
REASON_DUPLICATE_FILE = "aynı alan adı dosyada tekrar ediyor"
REASON_UNEXPECTED = "beklenmeyen hata"
WARN_BAD_EMAIL = "geçersiz e-posta biçimi, boş bırakıldı"

_DOMAIN_PATTERN = re.compile(r"[a-z0-9][a-z0-9-]*(\.[a-z0-9-]+)+")

SOURCE_CSV = "csv"
SOURCE_MANUAL = "manual"
SOURCE_TEKNOKENT = "teknokent"  # bkz. src/company_finder/teknokent_source.py (Faz 9.2)


class RowError(ValueError):
    """Satır (ya da elle girilen kayıt) kabul edilemez. Mesaj Türkçe ve kullanıcıya gösterilebilir."""


@dataclass(frozen=True)
class AddedRow:
    row: int
    name: str
    domain: str


@dataclass(frozen=True)
class RowNote:
    row: int
    name: str
    reason: str


@dataclass
class ImportReport:
    added: list[AddedRow] = field(default_factory=list)
    skipped: list[RowNote] = field(default_factory=list)
    warnings: list[RowNote] = field(default_factory=list)
    total_rows: int = 0

    @property
    def added_count(self) -> int:
        return len(self.added)

    @property
    def skipped_count(self) -> int:
        return len(self.skipped)

    def summary(self) -> str:
        parts = [f"{self.total_rows} satır okundu", f"{self.added_count} şirket eklendi"]
        if self.skipped_count:
            parts.append(f"{self.skipped_count} satır atlandı")
        if self.warnings:
            parts.append(f"{len(self.warnings)} uyarı")
        return ", ".join(parts) + "."


def _normalize_header(value: str) -> str:
    """Başlık hücresini eşleme için sadeleştirir: BOM, Türkçe katlama, tek boşluk."""
    cleaned = value.lstrip("﻿").strip()
    return re.sub(r"\s+", " ", fold_turkish(cleaned))


def _map_columns(header: list[str]) -> dict[str, int]:
    """Başlık satırını `{alan_adı: sütun_indeksi}`'ne çevirir; ilk eşleşen sütun kazanır."""
    normalized_alias_to_field = {
        _normalize_header(alias): canonical
        for canonical, aliases in COLUMN_ALIASES.items()
        for alias in aliases
    }
    mapping: dict[str, int] = {}
    for index, raw_cell in enumerate(header):
        canonical = normalized_alias_to_field.get(_normalize_header(raw_cell))
        if canonical is None:
            logger.debug(f"Tanınmayan sütun yok sayıldı: {raw_cell!r}")
            continue
        if canonical in mapping:
            logger.warning(f"'{canonical}' alanına birden çok sütun eşleşti, ilki kullanılıyor")
            continue
        mapping[canonical] = index
    return mapping


def _row_to_fields(row: list[str], column_map: dict[str, int]) -> dict[str, str]:
    """Ham hücre listesini, sütun eşlemesine göre kanonik alan sözlüğüne çevirir."""
    fields: dict[str, str] = {}
    for canonical, index in column_map.items():
        fields[canonical] = row[index].strip() if index < len(row) else ""
    return fields


def _decode(path: Path) -> str:
    """CSV dosyasını, sırayla `CSV_ENCODINGS` denenerek metne çevirir."""
    raw = path.read_bytes()
    if raw[:2] in _UTF16_BOMS:
        raise ValueError(
            "Dosya kodlaması desteklenmiyor (UTF-16 görünüyor). Dosyayı UTF-8 olarak kaydedin."
        )
    for encoding in CSV_ENCODINGS:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Dosya kodlaması okunamadı (utf-8 / utf-8-sig / cp1254 hiçbiri uymadı).")


def _pick_delimiter(header_line: str) -> str:
    """Başlık satırını `,` ve `;` ile ayrıştırıp, daha çok bilinen sütun tanıyanı seçer."""
    best_delimiter = ","
    best_score = (-1, -1)
    for delimiter in (",", ";"):
        cells = next(csv.reader(io.StringIO(header_line), delimiter=delimiter))
        mapped = len(_map_columns(cells))
        score = (mapped, len(cells))
        if score > best_score:
            best_score = score
            best_delimiter = delimiter
    return best_delimiter


def _read_csv_rows(path: Path) -> Iterator[tuple[int, list[str]]]:
    """CSV dosyasını satır satır döner: `(dosyadaki_satır_no, hücreler)`."""
    text = _decode(path)
    stream = io.StringIO(text, newline="")
    first_line = stream.readline()
    if not first_line:
        return
    delimiter = _pick_delimiter(first_line)
    stream.seek(0)
    reader = csv.reader(stream, delimiter=delimiter)
    for row in reader:
        if any(cell.strip() for cell in row):
            yield reader.line_num, row


def _read_xlsx_rows(path: Path) -> Iterator[tuple[int, list[str]]]:
    """XLSX dosyasını satır satır döner: `(sayfadaki_satır_no, hücreler)`."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError) as exc:
        raise ValueError(f"Excel dosyası okunamadı: {exc}") from exc
    try:
        sheet = workbook.active or workbook.worksheets[0]
        for row_number, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = ["" if cell is None else str(cell).strip() for cell in cells]
            if any(values):
                yield row_number, values
    finally:
        workbook.close()


def _read_rows(path: Path) -> Iterator[tuple[int, list[str]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv_rows(path)
    if suffix == ".xlsx":
        return _read_xlsx_rows(path)
    raise ValueError(f"Desteklenmeyen dosya türü: {suffix or '(uzantısız)'} (.csv veya .xlsx bekleniyor)")


def _clean_website(value: str) -> Optional[str]:
    """Siteyi sadeleştirir: boşsa None, şeması yoksa https:// eklenir."""
    cleaned = value.replace("\xa0", " ").strip()
    if not cleaned:
        return None
    if not re.match(r"^https?://", cleaned, re.IGNORECASE):
        cleaned = f"https://{cleaned}"
    return cleaned


def _resolve_domain(website: Optional[str], email: Optional[str]) -> str:
    """Siteden, yoksa e-postadan alan adı çıkarır ve biçimini doğrular.

    `normalize_domain()` bir doğrulayıcı değildir (bkz. modül docstring'i); bu yüzden sonuç
    burada `_DOMAIN_PATTERN` ile kontrol edilir.
    """
    raw = website if website else email.rpartition("@")[2] if email else ""
    domain = normalize_domain(raw)
    if not _DOMAIN_PATTERN.fullmatch(domain):
        raise RowError(REASON_BAD_DOMAIN)
    return domain


def _parse_tags(value: str) -> list[str]:
    """Etiket hücresini böler, Türkçe katlar, sözlüğe göre süzer (en fazla `MAX_TAGS`)."""
    raw_tags = re.split(r"[,;|]", value)
    slugs = [re.sub(r"[\s_]+", "-", fold_turkish(tag)).strip("-") for tag in raw_tags]
    return filter_tags(slugs)


def _build_company(
    fields: dict[str, str],
    *,
    source: str,
    db_domains: set[str],
    file_domains: set[str],
    blocked: frozenset[str] | set[str] = frozenset(),
) -> tuple[Company, Optional[str]]:
    """Alan sözlüğünden bir `Company` üretir. Geçersizse `RowError` fırlatır.

    Döndürülen ikilinin ikinci elemanı, satır kabul edildiyse ama bir uyarısı varsa
    (şu an sadece bozuk e-posta) Türkçe uyarı metnidir, yoksa `None`.

    `db_domains` (veritabanında zaten var) ve `file_domains` (bu çalıştırmada dosyada daha
    önce görüldü) ayrı tutulur ki rapor iki tekrar türünü birbirinden ayırt edebilsin.

    `blocked` (Adım 10.8b, `core/blocklist.py`): engelli alan adları; şirketin alan adı YA DA
    e-postasının alan adı bunlardan biriyse satır reddedilir. Bu, CSV dahil TÜM kaynakların
    (teknokent/OSM/Google `add_company()` üzerinden) ortak geçididir.
    """
    name = fields.get("name", "").strip()
    if not name:
        raise RowError(REASON_NO_NAME)

    warning: Optional[str] = None
    raw_email = fields.get("email", "")
    email = normalize_email(raw_email)
    if raw_email.strip() and email is None:
        warning = WARN_BAD_EMAIL

    website = _clean_website(fields.get("website", ""))

    if not website and not email:
        raise RowError(REASON_NO_CONTACT)

    domain = _resolve_domain(website, email)

    if is_free_mail_domain(domain):
        raise RowError(REASON_FREE_MAIL_NEEDS_WEBSITE)

    email_domain = normalize_domain(email.rpartition("@")[2]) if email else ""
    if domain in blocked or (email_domain and email_domain in blocked):
        raise RowError(REASON_BLOCKED)

    if domain in db_domains:
        raise RowError(REASON_DUPLICATE_DB)
    if domain in file_domains:
        raise RowError(REASON_DUPLICATE_FILE)

    company = Company(
        name=name,
        website=website,
        domain=domain,
        email=email,
        email_source=source if email else None,
        # CSV/elle girilen adres de rol tespitinden geçer: aksi halde CLAUDE.md 10. kural
        # (rol olmayan adrese otomatik gönderim yok) sadece taranan şirketlerde uygulanırdı.
        email_is_role=is_role_address(email) if email else None,
        city=fields.get("city", "").strip() or None,
        about_text=fields.get("about_text", "").strip() or None,
        sector_tags=_parse_tags(fields.get("tags", "")),
        source=source,
    )
    return company, warning


def _known_domains(session: Session) -> set[str]:
    """Veritabanındaki mevcut alan adlarını (normalize edilmiş) tek sorguda toplar."""
    return {normalize_domain(domain) for domain in session.exec(select(Company.domain))}


def import_companies(session: Session, path: str | Path, dry_run: bool = False) -> ImportReport:
    """CSV/XLSX dosyasındaki şirketleri oturuma ekler.

    COMMIT ETMEZ ("commit çağırana aittir", `src/core/db.py`). `dry_run=True` iken
    `session.add()` hiç çağrılmaz; rapor yine tam hesaplanır.

    Dosya düzeyinde bir sorun varsa (desteklenmeyen uzantı, ad sütunu bulunamadı, kodlama
    çözülemedi) `ValueError` fırlatılır. Satır düzeyindeki her sorun rapora yazılır, içe
    aktarmayı durdurmaz.
    """
    path = Path(path)
    rows = _read_rows(path)

    report = ImportReport()
    db_domains = _known_domains(session)
    blocked = blocked_domains(session)
    file_domains: set[str] = set()
    column_map: Optional[dict[str, int]] = None

    for row_number, cells in rows:
        if column_map is None:
            column_map = _map_columns(cells)
            if "name" not in column_map:
                raise ValueError(
                    "Dosyada şirket adı sütunu bulunamadı (ör. 'name', 'şirket adı')"
                )
            continue

        report.total_rows += 1
        fields = _row_to_fields(cells, column_map)
        name = fields.get("name", "").strip()

        try:
            company, warning = _build_company(
                fields,
                source=SOURCE_CSV,
                db_domains=db_domains,
                file_domains=file_domains,
                blocked=blocked,
            )
        except RowError as exc:
            report.skipped.append(RowNote(row=row_number, name=name, reason=str(exc)))
            continue
        except Exception:
            logger.exception(f"{row_number}. satırda beklenmeyen hata")
            report.skipped.append(RowNote(row=row_number, name=name, reason=REASON_UNEXPECTED))
            continue

        file_domains.add(company.domain)
        if warning:
            report.warnings.append(RowNote(row=row_number, name=name, reason=warning))
        if not dry_run:
            session.add(company)
        report.added.append(AddedRow(row=row_number, name=company.name, domain=company.domain))

    return report


def add_company(
    session: Session,
    name: str,
    website: Optional[str] = None,
    email: Optional[str] = None,
    city: Optional[str] = None,
    about_text: Optional[str] = None,
    source: str = SOURCE_MANUAL,
) -> Company:
    """Tek bir şirketi elle ekler; `import_companies` ile aynı kurallara uyar.

    `source` varsayılan olarak "manual" (arayüzdeki "elle ekle" formu); teknokent kaynağı
    (`src/company_finder/teknokent_source.py`, Faz 9.2) `SOURCE_TEKNOKENT` geçirerek aynı
    doğrulama/tekilleştirme çekirdeğini kullanır.

    Hata durumunda Türkçe mesajlı `ValueError` (`RowError`) fırlatır. COMMIT ETMEZ;
    çağıran `session.commit()` yapmalıdır. `session.flush()` çağrılır, böylece
    dönen `Company.id` doludur.
    """
    fields = {
        "name": name,
        "website": website or "",
        "email": email or "",
        "city": city or "",
        "about_text": about_text or "",
        "tags": "",
    }
    db_domains = _known_domains(session)
    company, warning = _build_company(
        fields,
        source=source,
        db_domains=db_domains,
        file_domains=set(),
        blocked=blocked_domains(session),
    )
    if warning:
        logger.warning(f"{name}: {warning}")
    session.add(company)
    session.flush()
    return company
